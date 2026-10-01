"""Ticket 07: PDF and DOCX extraction."""

from __future__ import annotations

import asyncio
import io

import docx
from httpx import ASGITransport, AsyncClient
from pypdf import PdfReader, PdfWriter

from app.main import create_app

from .fakes import FakeEmbedder
from .helpers import login


def make_pdf(pages: list[str]) -> bytes:
    """Build a minimal text PDF without a third-party writer."""
    font_num = 3 + 2 * len(pages)
    objects: dict[int, bytes] = {}
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode()
    for i, text in enumerate(pages):
        page_num, content_num = 3 + 2 * i, 4 + 2 * i
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 24 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects[page_num] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_num} 0 R >> >> /Contents {content_num} 0 R >>"
        ).encode()
        objects[content_num] = (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objects[font_num] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {font_num + 1}\n".encode() + b"0000000000 65535 f \n"
    for number in range(1, font_num + 1):
        out += f"{offsets[number]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {font_num + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    document = docx.Document()
    for kind, text in paragraphs:
        if kind == "heading":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_encrypted_pdf(text: str, password: str) -> bytes:
    reader = PdfReader(io.BytesIO(make_pdf([text])))
    writer = PdfWriter()
    writer.append(reader)
    writer.encrypt(password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


async def _wait_for_status(client, document_id: str) -> dict:
    body = {}
    for _ in range(60):
        body = (await client.get(f"/api/documents/{document_id}")).json()
        if body["status"] in ("ready", "failed"):
            return body
        await asyncio.sleep(0.05)
    return body


def _upload(files):
    return [("files", (name, content, ctype)) for name, content, ctype in files]


async def _upload_and_wait(client, session_id: str, name: str, content: bytes, ctype: str) -> dict:
    response = await client.post(
        f"/api/sessions/{session_id}/documents", files=_upload([(name, content, ctype)])
    )
    assert response.status_code == 202
    return await _wait_for_status(client, response.json()[0]["id"])


async def _session(client) -> dict:
    return (await client.post("/api/sessions", json={"title": "Docs"})).json()


async def test_text_pdf_reaches_ready_with_page_locators(client, database):
    await login(client)
    session = await _session(client)
    pdf = make_pdf(["First page about cells", "Second page about proteins"])
    document = await _upload_and_wait(client, session["id"], "paper.pdf", pdf, "application/pdf")
    assert document["status"] == "ready"
    assert document["chunk_count"] > 0

    chunks = await database.db.chunks.find({"document_id": document["id"]}).to_list(length=None)
    pages = {chunk["locator"].get("page") for chunk in chunks}
    assert pages == {1, 2}


async def test_docx_reaches_ready(client):
    await login(client)
    session = await _session(client)
    docx_bytes = make_docx([("heading", "Introduction"), ("para", "Word document body text.")])
    document = await _upload_and_wait(
        client,
        session["id"],
        "report.docx",
        docx_bytes,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert document["status"] == "ready"
    assert document["chunk_count"] > 0


async def test_scanned_pdf_fails_with_no_extractable_text(client):
    await login(client)
    session = await _session(client)
    document = await _upload_and_wait(client, session["id"], "scan.pdf", make_pdf([""]), "application/pdf")
    assert document["status"] == "failed"
    assert "no extractable text" in document["error"].lower()


async def test_password_protected_pdf_fails_with_a_reason(client):
    await login(client)
    session = await _session(client)
    encrypted = make_encrypted_pdf("secret", "hunter2")
    document = await _upload_and_wait(client, session["id"], "locked.pdf", encrypted, "application/pdf")
    assert document["status"] == "failed"
    assert "password" in document["error"].lower()


async def test_corrupt_pdf_fails_rather_than_hanging(client):
    await login(client)
    session = await _session(client)
    document = await _upload_and_wait(
        client, session["id"], "broken.pdf", b"%PDF-1.4 definitely not a real pdf", "application/pdf"
    )
    assert document["status"] == "failed"
    assert document["error"]


async def test_large_pdf_upload_does_not_block_other_requests(settings, database):
    # A slow extractor must not block the rest of the app: ingestion is a
    # background task, so another request completes while it runs.
    app = create_app(settings=settings, db=database, embedder=FakeEmbedder())
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="https://testserver") as local:
        await login(local)
        session = (await local.post("/api/sessions", json={"title": "Busy"})).json()
        pages = [f"Page {i} with a bit of text to extract" for i in range(1, 40)]
        response = await local.post(
            f"/api/sessions/{session['id']}/documents",
            files=_upload([("big.pdf", make_pdf(pages), "application/pdf")]),
        )
        assert response.status_code == 202
        # The app is still serving reads.
        assert (await local.get("/api/sessions")).status_code == 200
        document = await _wait_for_status(local, response.json()[0]["id"])
        assert document["status"] == "ready"
