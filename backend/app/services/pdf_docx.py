"""PDF and DOCX extractors (ticket 07).

Registered into the same extraction registry as the text formats. PDFs record a
``page`` locator; DOCX sections record the nearest heading. A PDF with no text
layer yields no blocks, which the registry turns into the "no extractable text"
failure rather than a silently empty document.
"""

from __future__ import annotations

import io

from .extraction import TextBlock, UnsupportedFileType, register_extractor


def _extract_pdf(content: bytes, filename: str) -> list[TextBlock]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise UnsupportedFileType("PDF is password-protected")
    except UnsupportedFileType:
        raise
    except Exception as exc:  # pypdf raises a family of parse errors
        raise UnsupportedFileType(f"Could not read PDF — it may be corrupt: {exc}") from exc

    blocks: list[TextBlock] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if text:
            blocks.append(TextBlock(text=text, locator={"page": page_number}))
    return blocks


def _extract_docx(content: bytes, filename: str) -> list[TextBlock]:
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise UnsupportedFileType(f"Could not read DOCX: {exc}") from exc

    blocks: list[TextBlock] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        section = "\n".join(buffer).strip()
        if section:
            locator = {"section": heading} if heading else {}
            blocks.append(TextBlock(text=section, locator=locator))

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style_name.lower().startswith("heading"):
            flush()
            heading = text or heading
            buffer = [text] if text else []
        elif text:
            buffer.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                buffer.append(" | ".join(cells))

    flush()
    return blocks


register_extractor("pdf", _extract_pdf)
register_extractor("docx", _extract_docx)
